"""Exportación de resultados a CSV y Excel (.xlsx)."""

from __future__ import annotations

import csv
from pathlib import Path

from .models import Contact, ContactStatus

# Encabezados de las columnas exportadas.
HEADERS = ["Nombre", "Teléfono", "Correo", "Estado"]

# Colores (hexadecimal) usados para resaltar el estado en el Excel.
STATUS_COLORS = {
    ContactStatus.VALIDO: "C6EFCE",  # verde
    ContactStatus.REVISAR_TELEFONO: "FFEB9C",  # amarillo
    ContactStatus.INCOMPLETO: "FFC7CE",  # rojo
}


def _rows(contacts: list[Contact]) -> list[list[str]]:
    """Convierte los contactos en filas de texto plano para exportar."""
    return [[c.name, c.phone or "", c.email or "", c.status.value] for c in contacts]


def export_csv(path: Path | str, contacts: list[Contact]) -> None:
    """Exporta a CSV con codificación utf-8-sig (compatible con Excel)."""
    path = Path(path)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(HEADERS)
        writer.writerows(_rows(contacts))


def export_xlsx(path: Path | str, contacts: list[Contact]) -> None:
    """Exporta a Excel (.xlsx) con colores por estado y anchos ajustados."""
    path = Path(path)

    # Importación local para no bloquear el arranque si falta el paquete.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Contactos"

    # Encabezado con fondo azul y texto en blanco.
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    # Datos.
    for contact in contacts:
        ws.append([contact.name, contact.phone or "", contact.email or "", contact.status.value])

    # Colorear la columna "Estado" según el valor.
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        cell = row[0]
        status = ContactStatus(cell.value) if cell.value else ContactStatus.INCOMPLETO
        cell.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(status, "FFFFFF"))

    # Ajustar el ancho de cada columna a su contenido.
    for idx in range(1, len(HEADERS) + 1):
        letter = get_column_letter(idx)
        longest = 0
        for row in ws.iter_rows(min_row=1, min_col=idx, max_col=idx):
            value = row[0].value
            if value:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 45)

    wb.save(path)
